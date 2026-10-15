#!/usr/bin/env python3
"""Inventory SMHI official historical land-warning raw data without assuming its schema.

Source: SMHI 'Uppföljning av vädervarningar', 2021-10-26 to 2025-12-31.
This is national LAND warning raw data, not marine warnings or confirmed Luleå intersections.
"""
from __future__ import annotations
import io,json,os,hashlib,re
from datetime import datetime,timezone
from pathlib import Path
import requests
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/"data"/"warning_history_smhi"
OUT=ROOT/"docs"/"warnings"
SOURCE_PAGE="https://www.smhi.se/kunskapsbanken/meteorologi/varningar-och-meddelanden/uppfoljning-av-vadervarningar"
SOURCE_XLSX="https://www.smhi.se/download/18.3b0c31d719c28a35067254a/1771572701001/Konsekvensbaserade%20v%C3%A4dervarningar%2026%20okt%202021%20%2031%20dec%202025.xlsx"
def main():
    response=requests.get(SOURCE_XLSX,timeout=120,headers={"User-Agent":"Lulea-statistik historical warning research/1.0"})
    response.raise_for_status()
    raw=response.content
    if raw[:2]!=b"PK":raise RuntimeError("SMHI historical workbook is not an XLSX ZIP.")
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    overview=[]
    for ws in wb:
        first=[]
        for i,row in enumerate(ws.iter_rows(values_only=True)):
            nonempty=[{"column":j+1,"value":str(v)[:180]} for j,v in enumerate(row) if v is not None]
            if nonempty:first.append({"row":i+1,"cells":nonempty[:35]})
            if len(first)>=8:break
        overview.append({"sheet":ws.title,"rows":ws.max_row,"columns":ws.max_column,"first_nonempty_rows":first})
    BASE.mkdir(parents=True,exist_ok=True);OUT.mkdir(parents=True,exist_ok=True)
    (BASE/"official_land_warnings_2021_2025.xlsx").write_bytes(raw)
    generated=datetime.now(timezone.utc).isoformat()
    metadata={"generated_utc":generated,"source_page":SOURCE_PAGE,"source_xlsx":SOURCE_XLSX,"sha256":hashlib.sha256(raw).hexdigest(),
       "period_start":"2021-10-26","period_end":"2025-12-31","scope":"National land weather warnings only",
       "limitations":["Not marine warnings","Raw published warnings may include replacements and cancellations",
          "No municipality assignment until the geographic columns and geometries have been inspected"],
       "sheets":overview}
    (OUT/"historical_inventory.json").write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"sheets":[{"sheet":x["sheet"],"rows":x["rows"],"columns":x["columns"]} for x in overview],"sha256":metadata["sha256"]},ensure_ascii=False))
if __name__=="__main__":main()
