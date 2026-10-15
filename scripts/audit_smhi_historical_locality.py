#!/usr/bin/env python3
"""Conservative locality audit of SMHI historical LAND warnings.

No polygons in workbook. Never assert that a Norrbotten region label proves a
geographic intersection with Luleå municipality.
"""
import csv,json,re
from collections import Counter
from datetime import datetime,date,timezone
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"data/warning_history_smhi/official_land_warnings_2021_2025.xlsx"
OUT=ROOT/"docs/warnings"
REGION_COLUMNS=("Huvudvarningen Områdesnamn","Delvarningen Områdesnamn","Berörda områden","Delvarning Var")
DIRECT=re.compile(r"\bluleå\b|\blulea\b",re.IGNORECASE)
REGIONAL=re.compile(r"norrbotten|bottenviken|norra norrland|södra norrland|nordligaste delen av sverige",re.IGNORECASE)

def printable(value):
    if value is None:return ""
    if isinstance(value,(datetime,date)):return value.isoformat(sep=" ") if isinstance(value,datetime) else value.isoformat()
    return str(value).strip()

def main():
    if not SOURCE.exists():raise RuntimeError("Missing source workbook; run historical import first")
    wb=load_workbook(SOURCE,read_only=True,data_only=True)
    if "Rådata" not in wb:raise RuntimeError("Source workbook lacks Rådata")
    sheet=wb["Rådata"]
    lines=sheet.iter_rows(values_only=True)
    header=[printable(x) for x in next(lines)]
    needed={"Varnings ID","Första gällande datum","Varningstyp","Varningsnivå",*REGION_COLUMNS}
    if not needed.issubset(header):raise RuntimeError("Unexpected source columns: "+", ".join(sorted(needed-set(header))))
    records=[]
    for index,row in enumerate(lines,start=2):
        d=dict(zip(header,[printable(v) for v in row]))
        if not any(d.values()):continue
        geo=" ".join(d.get(k,"") for k in REGION_COLUMNS)
        if DIRECT.search(geo):status="explicit_lulea"
        elif REGIONAL.search(geo):status="regional_candidate"
        else:status="not_identified"
        records.append({"source_row":index,"warning_id":d.get("Varnings ID"),"date":d.get("Första gällande datum"),
            "type":d.get("Varningstyp"),"description":d.get("Varningsbeskrivning"),
            "level":d.get("Varningsnivå"),"start":d.get("Starttid"),"end":d.get("Sluttid"),
            "published":d.get("Publicerad"),"unpublished":d.get("Avpublicerad"),
            "main_area":d.get("Huvudvarningen Områdesnamn"),
            "sub_area":d.get("Delvarningen Områdesnamn"),
            "affected_areas":d.get("Berörda områden"),"location_text":d.get("Delvarning Var"),
            "classification":status})
    OUT.mkdir(parents=True,exist_ok=True)
    total=Counter(r["classification"] for r in records)
    summary={"generated_utc":datetime.now(timezone.utc).isoformat(),"source":"SMHI official workbook, Rådata",
        "total_records":len(records),"classifications":dict(total),
        "interpretation":{"explicit_lulea":"Luleå named in a geographic field; not confirmed polygon intersection",
                          "regional_candidate":"Broad regional wording may include Luleå; needs manual verification or historic geometry",
                          "not_identified":"No geographic indication matching the regional search; not proof Luleå is excluded"},
        "limitations":["Land weather warnings only","Textual candidate identification, not spatial intersection",
            "Source rows need not be unique meteorological events; preserve warning IDs and publication dates"]}
    (OUT/"historical_locality_audit.json").write_text(json.dumps({"summary":summary,"records":records},ensure_ascii=False,separators=(",",":")),encoding="utf-8")
    (OUT/"historical_locality_summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(summary,ensure_ascii=False))
if __name__=="__main__":main()
